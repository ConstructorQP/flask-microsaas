from flask import Flask, render_template, request, redirect, url_for, session, send_file
import sqlite3, io, os
from functools import wraps

app = Flask(__name__)
app.secret_key = 'constructor_microsaas_secret_2024'
DB_PATH = 'database.db'

# ─── DB ──────────────────────────────────────────────────────────────────────

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS usuarios (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password TEXT NOT NULL)''')
    c.execute('''CREATE TABLE IF NOT EXISTS clientes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre TEXT NOT NULL, email TEXT, telefono TEXT, direccion TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''')
    c.execute('''CREATE TABLE IF NOT EXISTS productos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre TEXT NOT NULL, descripcion TEXT, precio REAL NOT NULL DEFAULT 0, stock INTEGER DEFAULT 0,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''')
    c.execute('''CREATE TABLE IF NOT EXISTS pedidos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        cliente_id INTEGER NOT NULL, estado TEXT DEFAULT 'pendiente', total REAL DEFAULT 0, notas TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (cliente_id) REFERENCES clientes(id))''')
    c.execute('''CREATE TABLE IF NOT EXISTS lineas_pedido (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        pedido_id INTEGER NOT NULL, producto_id INTEGER NOT NULL,
        cantidad INTEGER NOT NULL DEFAULT 1, precio_unitario REAL NOT NULL, subtotal REAL NOT NULL,
        FOREIGN KEY (pedido_id) REFERENCES pedidos(id),
        FOREIGN KEY (producto_id) REFERENCES productos(id))''')
    try:
        c.execute("INSERT INTO usuarios (username, password) VALUES (?,?)", ('constructor', 'tryme'))
    except sqlite3.IntegrityError:
        pass
    conn.commit()
    conn.close()

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated

# ─── AUTH ────────────────────────────────────────────────────────────────────

@app.route('/', methods=['GET', 'POST'])
def login():
    if 'user' in session:
        return redirect(url_for('dashboard'))
    error = None
    if request.method == 'POST':
        u, p = request.form['username'], request.form['password']
        conn = get_db()
        user = conn.execute("SELECT * FROM usuarios WHERE username=? AND password=?", (u, p)).fetchone()
        conn.close()
        if user:
            session['user'] = u
            return redirect(url_for('dashboard'))
        error = 'Usuario o contraseña incorrectos'
    return render_template('login.html', error=error)

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

# ─── DASHBOARD ───────────────────────────────────────────────────────────────

@app.route('/dashboard')
@login_required
def dashboard():
    conn = get_db()
    n_clientes  = conn.execute("SELECT COUNT(*) FROM clientes").fetchone()[0]
    n_productos = conn.execute("SELECT COUNT(*) FROM productos").fetchone()[0]
    n_pedidos   = conn.execute("SELECT COUNT(*) FROM pedidos").fetchone()[0]
    total_ventas = conn.execute("SELECT COALESCE(SUM(total),0) FROM pedidos").fetchone()[0]
    conn.close()
    return render_template('dashboard.html', active='dashboard',
        n_clientes=n_clientes, n_productos=n_productos,
        n_pedidos=n_pedidos, total_ventas=total_ventas)

# ─── CLIENTES ────────────────────────────────────────────────────────────────

@app.route('/clientes')
@login_required
def clientes():
    conn = get_db()
    rows = conn.execute("SELECT * FROM clientes ORDER BY created_at DESC").fetchall()
    conn.close()
    return render_template('clientes.html', active='clientes', clientes=rows)

@app.route('/clientes/nuevo', methods=['GET', 'POST'])
@login_required
def cliente_nuevo():
    if request.method == 'POST':
        conn = get_db()
        conn.execute("INSERT INTO clientes (nombre,email,telefono,direccion) VALUES (?,?,?,?)",
            (request.form['nombre'], request.form['email'], request.form['telefono'], request.form['direccion']))
        conn.commit(); conn.close()
        return redirect(url_for('clientes'))
    return render_template('cliente_form.html', active='clientes', cliente=None)

@app.route('/clientes/editar/<int:id>', methods=['GET', 'POST'])
@login_required
def cliente_editar(id):
    conn = get_db()
    if request.method == 'POST':
        conn.execute("UPDATE clientes SET nombre=?,email=?,telefono=?,direccion=? WHERE id=?",
            (request.form['nombre'], request.form['email'], request.form['telefono'], request.form['direccion'], id))
        conn.commit(); conn.close()
        return redirect(url_for('clientes'))
    c = conn.execute("SELECT * FROM clientes WHERE id=?", (id,)).fetchone()
    conn.close()
    return render_template('cliente_form.html', active='clientes', cliente=c)

@app.route('/clientes/eliminar/<int:id>')
@login_required
def cliente_eliminar(id):
    conn = get_db()
    conn.execute("DELETE FROM clientes WHERE id=?", (id,))
    conn.commit(); conn.close()
    return redirect(url_for('clientes'))

# ─── CLIENTES EXPORT ─────────────────────────────────────────────────────────

@app.route('/clientes/exportar/excel')
@login_required
def clientes_exportar_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    conn = get_db()
    rows = conn.execute("SELECT * FROM clientes ORDER BY created_at DESC").fetchall()
    conn.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    headers = ['ID', 'Nombre', 'Email', 'Teléfono', 'Dirección', 'Fecha alta']
    header_fill = PatternFill('solid', start_color='3730A3')
    header_font = Font(bold=True, color='FFFFFF', name='Arial', size=10)
    thin = Side(style='thin', color='D1D5DB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    ws.row_dimensions[1].height = 22
    for row_idx, row in enumerate(rows, 2):
        vals = [row['id'], row['nombre'], row['email'] or '', row['telefono'] or '', row['direccion'] or '', str(row['created_at'])[:10]]
        fill = PatternFill('solid', start_color='EEF2FF') if row_idx % 2 == 0 else PatternFill('solid', start_color='FFFFFF')
        for col, v in enumerate(vals, 1):
            cell = ws.cell(row=row_idx, column=col, value=v)
            cell.font = Font(name='Arial', size=10)
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(vertical='center')
    for col, w in enumerate([8, 30, 30, 16, 40, 14], 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w
    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='clientes.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/clientes/exportar/pdf')
@login_required
def clientes_exportar_pdf():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    conn = get_db()
    rows = conn.execute("SELECT * FROM clientes ORDER BY created_at DESC").fetchall()
    conn.close()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=30, bottomMargin=30, leftMargin=30, rightMargin=30)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('title', fontName='Helvetica-Bold', fontSize=16, textColor=colors.HexColor('#3730A3'), spaceAfter=8)
    sub_style   = ParagraphStyle('sub',   fontName='Helvetica', fontSize=9, textColor=colors.HexColor('#6B7280'), spaceAfter=16)
    story = [
        Paragraph('constructor | microsaas', title_style),
        Paragraph('Listado de Clientes', sub_style),
    ]
    data = [['ID', 'Nombre', 'Email', 'Teléfono', 'Dirección', 'Fecha alta']]
    for r in rows:
        data.append([str(r['id']), r['nombre'], r['email'] or '', r['telefono'] or '', r['direccion'] or '', str(r['created_at'])[:10]])
    t = Table(data, colWidths=[30, 130, 150, 80, 180, 70])
    t.setStyle(TableStyle([
        ('BACKGROUND',   (0,0), (-1,0), colors.HexColor('#3730A3')),
        ('TEXTCOLOR',    (0,0), (-1,0), colors.white),
        ('FONTNAME',     (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE',     (0,0), (-1,0), 9),
        ('FONTNAME',     (0,1), (-1,-1), 'Helvetica'),
        ('FONTSIZE',     (0,1), (-1,-1), 8),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor('#EEF2FF'), colors.white]),
        ('GRID',         (0,0), (-1,-1), 0.4, colors.HexColor('#D1D5DB')),
        ('VALIGN',       (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING',   (0,0), (-1,-1), 5),
        ('BOTTOMPADDING',(0,0), (-1,-1), 5),
    ]))
    story.append(t)
    doc.build(story)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='clientes.pdf', mimetype='application/pdf')

# ─── PRODUCTOS ───────────────────────────────────────────────────────────────

@app.route('/productos')
@login_required
def productos():
    conn = get_db()
    rows = conn.execute("SELECT * FROM productos ORDER BY created_at DESC").fetchall()
    conn.close()
    return render_template('productos.html', active='productos', productos=rows)

@app.route('/productos/nuevo', methods=['GET', 'POST'])
@login_required
def producto_nuevo():
    if request.method == 'POST':
        conn = get_db()
        conn.execute("INSERT INTO productos (nombre,descripcion,precio,stock) VALUES (?,?,?,?)",
            (request.form['nombre'], request.form['descripcion'], request.form['precio'], request.form['stock']))
        conn.commit(); conn.close()
        return redirect(url_for('productos'))
    return render_template('producto_form.html', active='productos', producto=None)

@app.route('/productos/editar/<int:id>', methods=['GET', 'POST'])
@login_required
def producto_editar(id):
    conn = get_db()
    if request.method == 'POST':
        conn.execute("UPDATE productos SET nombre=?,descripcion=?,precio=?,stock=? WHERE id=?",
            (request.form['nombre'], request.form['descripcion'], request.form['precio'], request.form['stock'], id))
        conn.commit(); conn.close()
        return redirect(url_for('productos'))
    p = conn.execute("SELECT * FROM productos WHERE id=?", (id,)).fetchone()
    conn.close()
    return render_template('producto_form.html', active='productos', producto=p)

@app.route('/productos/eliminar/<int:id>')
@login_required
def producto_eliminar(id):
    conn = get_db()
    conn.execute("DELETE FROM productos WHERE id=?", (id,))
    conn.commit(); conn.close()
    return redirect(url_for('productos'))

@app.route('/productos/exportar/excel')
@login_required
def productos_exportar_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    conn = get_db()
    rows = conn.execute("SELECT * FROM productos ORDER BY created_at DESC").fetchall()
    conn.close()
    wb = Workbook(); ws = wb.active; ws.title = "Productos"
    headers = ['ID', 'Nombre', 'Descripción', 'Precio (€)', 'Stock', 'Fecha alta']
    header_fill = PatternFill('solid', start_color='EA580C')
    thin = Side(style='thin', color='D1D5DB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color='FFFFFF', name='Arial', size=10)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    ws.row_dimensions[1].height = 22
    for ri, row in enumerate(rows, 2):
        fill = PatternFill('solid', start_color='FFF7ED') if ri % 2 == 0 else PatternFill('solid', start_color='FFFFFF')
        vals = [row['id'], row['nombre'], row['descripcion'] or '', row['precio'], row['stock'], str(row['created_at'])[:10]]
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.font = Font(name='Arial', size=10)
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(vertical='center')
    for col, w in enumerate([8, 30, 40, 14, 12, 14], 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='productos.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/productos/exportar/pdf')
@login_required
def productos_exportar_pdf():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    conn = get_db()
    rows = conn.execute("SELECT * FROM productos ORDER BY created_at DESC").fetchall()
    conn.close()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=30, bottomMargin=30, leftMargin=30, rightMargin=30)
    title_s = ParagraphStyle('t', fontName='Helvetica-Bold', fontSize=16, textColor=colors.HexColor('#EA580C'), spaceAfter=8)
    sub_s   = ParagraphStyle('s', fontName='Helvetica', fontSize=9,  textColor=colors.HexColor('#6B7280'), spaceAfter=16)
    story = [Paragraph('constructor | microsaas', title_s), Paragraph('Listado de Productos', sub_s)]
    data = [['ID', 'Nombre', 'Descripción', 'Precio (€)', 'Stock', 'Fecha alta']]
    for r in rows:
        data.append([str(r['id']), r['nombre'], r['descripcion'] or '', f"{r['precio']:.2f}", str(r['stock']), str(r['created_at'])[:10]])
    t = Table(data, colWidths=[30, 130, 200, 70, 50, 70])
    t.setStyle(TableStyle([
        ('BACKGROUND',   (0,0), (-1,0), colors.HexColor('#EA580C')),
        ('TEXTCOLOR',    (0,0), (-1,0), colors.white),
        ('FONTNAME',     (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE',     (0,0), (-1,0), 9),
        ('FONTNAME',     (0,1), (-1,-1), 'Helvetica'),
        ('FONTSIZE',     (0,1), (-1,-1), 8),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor('#FFF7ED'), colors.white]),
        ('GRID',         (0,0), (-1,-1), 0.4, colors.HexColor('#D1D5DB')),
        ('VALIGN',       (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING',   (0,0), (-1,-1), 5),
        ('BOTTOMPADDING',(0,0), (-1,-1), 5),
    ]))
    story.append(t)
    doc.build(story)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='productos.pdf', mimetype='application/pdf')

# ─── PEDIDOS ─────────────────────────────────────────────────────────────────

@app.route('/pedidos')
@login_required
def pedidos():
    conn = get_db()
    rows = conn.execute("""SELECT p.*, c.nombre as cliente_nombre
        FROM pedidos p JOIN clientes c ON p.cliente_id=c.id
        ORDER BY p.created_at DESC""").fetchall()
    conn.close()
    return render_template('pedidos.html', active='pedidos', pedidos=rows)

@app.route('/pedidos/nuevo', methods=['GET', 'POST'])
@login_required
def pedido_nuevo():
    conn = get_db()
    if request.method == 'POST':
        cur = conn.execute("INSERT INTO pedidos (cliente_id,estado,notas,total) VALUES (?,?,?,0)",
            (request.form['cliente_id'], request.form['estado'], request.form['notas']))
        pedido_id = cur.lastrowid
        total = 0
        for pid, qty in zip(request.form.getlist('producto_id[]'), request.form.getlist('cantidad[]')):
            if pid and qty and int(qty) > 0:
                prod = conn.execute("SELECT precio FROM productos WHERE id=?", (pid,)).fetchone()
                if prod:
                    sub = prod['precio'] * int(qty)
                    total += sub
                    conn.execute("INSERT INTO lineas_pedido (pedido_id,producto_id,cantidad,precio_unitario,subtotal) VALUES (?,?,?,?,?)",
                        (pedido_id, pid, qty, prod['precio'], sub))
        conn.execute("UPDATE pedidos SET total=? WHERE id=?", (total, pedido_id))
        conn.commit(); conn.close()
        return redirect(url_for('pedidos'))
    clientes = conn.execute("SELECT * FROM clientes ORDER BY nombre").fetchall()
    prods    = conn.execute("SELECT * FROM productos ORDER BY nombre").fetchall()
    conn.close()
    return render_template('pedido_form.html', active='pedidos', pedido=None, clientes=clientes, productos=prods, lineas=[])

@app.route('/pedidos/ver/<int:id>')
@login_required
def pedido_ver(id):
    conn = get_db()
    pedido = conn.execute("""SELECT p.*, c.nombre as cliente_nombre, c.email as cliente_email
        FROM pedidos p JOIN clientes c ON p.cliente_id=c.id WHERE p.id=?""", (id,)).fetchone()
    lineas = conn.execute("""SELECT lp.*, pr.nombre as producto_nombre
        FROM lineas_pedido lp JOIN productos pr ON lp.producto_id=pr.id WHERE lp.pedido_id=?""", (id,)).fetchall()
    conn.close()
    return render_template('pedido_ver.html', active='pedidos', pedido=pedido, lineas=lineas)

@app.route('/pedidos/editar/<int:id>', methods=['GET', 'POST'])
@login_required
def pedido_editar(id):
    conn = get_db()
    if request.method == 'POST':
        conn.execute("DELETE FROM lineas_pedido WHERE pedido_id=?", (id,))
        total = 0
        for pid, qty in zip(request.form.getlist('producto_id[]'), request.form.getlist('cantidad[]')):
            if pid and qty and int(qty) > 0:
                prod = conn.execute("SELECT precio FROM productos WHERE id=?", (pid,)).fetchone()
                if prod:
                    sub = prod['precio'] * int(qty)
                    total += sub
                    conn.execute("INSERT INTO lineas_pedido (pedido_id,producto_id,cantidad,precio_unitario,subtotal) VALUES (?,?,?,?,?)",
                        (id, pid, qty, prod['precio'], sub))
        conn.execute("UPDATE pedidos SET cliente_id=?,estado=?,notas=?,total=? WHERE id=?",
            (request.form['cliente_id'], request.form['estado'], request.form['notas'], total, id))
        conn.commit(); conn.close()
        return redirect(url_for('pedidos'))
    pedido   = conn.execute("SELECT * FROM pedidos WHERE id=?", (id,)).fetchone()
    clientes = conn.execute("SELECT * FROM clientes ORDER BY nombre").fetchall()
    prods    = conn.execute("SELECT * FROM productos ORDER BY nombre").fetchall()
    lineas   = conn.execute("""SELECT lp.*, pr.nombre as producto_nombre
        FROM lineas_pedido lp JOIN productos pr ON lp.producto_id=pr.id WHERE lp.pedido_id=?""", (id,)).fetchall()
    conn.close()
    return render_template('pedido_form.html', active='pedidos', pedido=pedido, clientes=clientes, productos=prods, lineas=lineas)

@app.route('/pedidos/eliminar/<int:id>')
@login_required
def pedido_eliminar(id):
    conn = get_db()
    conn.execute("DELETE FROM lineas_pedido WHERE pedido_id=?", (id,))
    conn.execute("DELETE FROM pedidos WHERE id=?", (id,))
    conn.commit(); conn.close()
    return redirect(url_for('pedidos'))

@app.route('/pedidos/exportar/excel')
@login_required
def pedidos_exportar_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    conn = get_db()
    rows = conn.execute("""SELECT p.*, c.nombre as cliente_nombre
        FROM pedidos p JOIN clientes c ON p.cliente_id=c.id ORDER BY p.created_at DESC""").fetchall()
    conn.close()
    wb = Workbook(); ws = wb.active; ws.title = "Pedidos"
    headers = ['Referencia', 'Cliente', 'Estado', 'Total (€)', 'Notas', 'Fecha']
    header_fill = PatternFill('solid', start_color='3730A3')
    thin = Side(style='thin', color='D1D5DB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color='FFFFFF', name='Arial', size=10)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    ws.row_dimensions[1].height = 22
    for ri, row in enumerate(rows, 2):
        fill = PatternFill('solid', start_color='EEF2FF') if ri % 2 == 0 else PatternFill('solid', start_color='FFFFFF')
        vals = [f"PED-{row['id']:04d}", row['cliente_nombre'], row['estado'].capitalize(), row['total'], row['notas'] or '', str(row['created_at'])[:10]]
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.font = Font(name='Arial', size=10)
            cell.fill = fill; cell.border = border
            cell.alignment = Alignment(vertical='center')
    for col, w in enumerate([14, 30, 14, 14, 40, 14], 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='pedidos.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/pedidos/exportar/pdf')
@login_required
def pedidos_exportar_pdf():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    conn = get_db()
    rows = conn.execute("""SELECT p.*, c.nombre as cliente_nombre
        FROM pedidos p JOIN clientes c ON p.cliente_id=c.id ORDER BY p.created_at DESC""").fetchall()
    conn.close()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=30, bottomMargin=30, leftMargin=30, rightMargin=30)
    title_s = ParagraphStyle('t', fontName='Helvetica-Bold', fontSize=16, textColor=colors.HexColor('#3730A3'), spaceAfter=8)
    sub_s   = ParagraphStyle('s', fontName='Helvetica', fontSize=9,  textColor=colors.HexColor('#6B7280'), spaceAfter=16)
    story = [Paragraph('constructor | microsaas', title_s), Paragraph('Listado de Pedidos', sub_s)]
    data = [['Referencia', 'Cliente', 'Estado', 'Total (€)', 'Notas', 'Fecha']]
    for r in rows:
        data.append([f"PED-{r['id']:04d}", r['cliente_nombre'], r['estado'].capitalize(),
                     f"{r['total']:.2f}", r['notas'] or '', str(r['created_at'])[:10]])
    t = Table(data, colWidths=[70, 130, 70, 70, 180, 70])
    t.setStyle(TableStyle([
        ('BACKGROUND',   (0,0), (-1,0), colors.HexColor('#3730A3')),
        ('TEXTCOLOR',    (0,0), (-1,0), colors.white),
        ('FONTNAME',     (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE',     (0,0), (-1,0), 9),
        ('FONTNAME',     (0,1), (-1,-1), 'Helvetica'),
        ('FONTSIZE',     (0,1), (-1,-1), 8),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor('#EEF2FF'), colors.white]),
        ('GRID',         (0,0), (-1,-1), 0.4, colors.HexColor('#D1D5DB')),
        ('VALIGN',       (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING',   (0,0), (-1,-1), 5),
        ('BOTTOMPADDING',(0,0), (-1,-1), 5),
    ]))
    story.append(t)
    doc.build(story)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='pedidos.pdf', mimetype='application/pdf')

if __name__ == '__main__':
    init_db()
    app.run(debug=True)
